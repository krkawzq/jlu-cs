"""Capture the actual web slides, then place those screenshots into a PPTX."""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import datetime, timezone
import hashlib
import json
import mmap
import os
from pathlib import Path
import shutil
import tempfile
import threading

from server import Application, PresentationServer, ROOT, default_runtime

BROWSERS = ROOT / 'build/presentation-tools/browsers'
if BROWSERS.is_dir():
    os.environ.setdefault('PLAYWRIGHT_BROWSERS_PATH', str(BROWSERS))

from PIL import Image, ImageDraw
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches


@contextmanager
def launch_browser(playwright):
    executable = Path(playwright.chromium.executable_path)
    temporary = None
    try:
        # Some shared mounts permit reads but not the mmap required by Chromium.
        try:
            with (executable.parent / 'icudtl.dat').open('rb') as stream:
                with mmap.mmap(stream.fileno(), 0, access=mmap.ACCESS_READ):
                    pass
        except OSError as error:
            if error.errno not in (19, 22, 95):
                raise
            temporary = tempfile.TemporaryDirectory(prefix='prim-presentation-browser-')
            local = Path(temporary.name) / 'browser'
            shutil.copytree(executable.parent, local)
            executable = local / executable.name
        browser = playwright.chromium.launch(executable_path=str(executable), headless=True, args=['--disable-gpu'])
        try:
            yield browser
        finally:
            browser.close()
    finally:
        if temporary:
            temporary.cleanup()


@contextmanager
def preview_server(outline, presenter='', course='', date='', runtime=None):
    app = Application(outline, runtime, presenter, course, date)
    server = PresentationServer(('127.0.0.1', 0), app)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f'http://127.0.0.1:{server.server_port}', app
    finally:
        app.runner.cancel()
        server.shutdown()
        server.server_close()
        thread.join(timeout=2)


def inspect_layout(page):
    return page.evaluate('''() => {
      const slide = document.querySelector('.slide.active');
      const frame = slide.getBoundingClientRect();
      const failures = [];
      for (const element of slide.querySelectorAll('h1,h2,p,table,.concept-card,.code-panel,.code-text,.pipe-stage,.evidence-card,svg,.slide-foot,figcaption')) {
        if (!element.getClientRects().length) continue;
        const r = element.getBoundingClientRect();
        if (r.top < frame.top - 1 || r.left < frame.left - 1 || r.bottom > frame.bottom + 1 || r.right > frame.right + 1)
          failures.push({kind:'outside-slide', tag:element.tagName,text:element.textContent.slice(0,80)});
        if (element.scrollWidth > element.clientWidth + 3 && element.tagName !== 'svg')
          failures.push({kind:'horizontal-overflow',tag:element.tagName,text:element.textContent.slice(0,80)});
      }
      const copy = slide.querySelector('.copy');
      const footer = slide.querySelector('.slide-foot');
      if (copy && copy.getBoundingClientRect().bottom > footer.getBoundingClientRect().top)
        failures.push({kind:'body-footer-overlap'});
      return failures;
    }''')


def capture(url, output: Path, main_only=False):
    output.parent.mkdir(parents=True, exist_ok=True)
    assets = output.parent / (output.stem + '-assets')
    assets.mkdir(exist_ok=True)
    errors = []
    with sync_playwright() as playwright, launch_browser(playwright) as browser:
        try:
            page = browser.new_page(viewport={'width':1600, 'height':900}, device_scale_factor=1,
                                    reduced_motion='reduce')
            page.on('pageerror', lambda error: errors.append(str(error)))
            page.goto(url.rstrip('/') + '/?export=1', wait_until='networkidle')
            page.wait_for_function('() => window.PrimPresentation?.ready === true')
            deck = page.evaluate('window.PrimPresentation.data')
            evidence = page.evaluate('window.PrimPresentation.evidence')
            slides = [slide for slide in deck['slides'] if not main_only or not slide['id'].startswith('A')]
            receipts = []
            for slide in slides:
                page.evaluate('(id) => window.PrimPresentation.show(id)', slide['id'])
                page.evaluate('() => new Promise(resolve => requestAnimationFrame(() => requestAnimationFrame(resolve)))')
                problems = inspect_layout(page)
                if problems:
                    raise ValueError(f"{slide['id']} 页面越界：{problems}")
                file = assets / f"slide-{slide['id']}.png"
                page.locator('.slide.active').screenshot(path=str(file))
                receipts.append({'id':slide['id'],'image':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest()})
                print(f"截图 {slide['id']}：{slide['title']}", flush=True)
            if errors:
                raise ValueError(f'浏览器错误：{errors}')
        finally:
            browser.close()

    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.title = deck['metadata']['title']
    presentation.core_properties.subject = '网页逐页截图；交互与代码执行请使用在线展示'
    presentation.core_properties.author = deck['presenter'].get('name') or 'Prim'
    for slide, receipt in zip(slides, receipts):
        target = presentation.slides.add_slide(presentation.slide_layouts[6])
        target.shapes.add_picture(str(assets / receipt['image']), 0, 0,
                                 width=presentation.slide_width, height=presentation.slide_height)
    temporary = output.with_suffix('.partial.pptx')
    presentation.save(temporary)
    temporary.replace(output)

    thumb_w, thumb_h, label_h = 400, 225, 25
    sheet = Image.new('RGB', (thumb_w * 3, (thumb_h + label_h) * ((len(slides) + 2)//3)), '#F0EEE6')
    draw = ImageDraw.Draw(sheet)
    for index, receipt in enumerate(receipts):
        x, y = index % 3 * thumb_w, index // 3 * (thumb_h + label_h)
        with Image.open(assets / receipt['image']) as image:
            sheet.paste(image.convert('RGB').resize((thumb_w,thumb_h),Image.Resampling.LANCZOS), (x,y))
        draw.text((x+12,y+thumb_h+5), receipt['id'], fill='#141413')
    sheet.save(assets / 'contact-sheet.jpg', quality=90)
    (assets / 'manifest.json').write_text(json.dumps({
        'generated_at':datetime.now(timezone.utc).isoformat(), 'source_sha256':deck['source_sha256'],
        'viewport':[1600,900], 'slides':receipts,
        'measurement': {'generated_at': evidence.get('generated_at'), 'runtime_sha256': evidence.get('runtime_sha256'), 'source_sha256': evidence.get('source_sha256')} if evidence else None,
        'pptx_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),
        'renderer':'Playwright Chromium; screenshots of actual HTML slides',
        'notes':'Slide images are static; every stepped element is shown.'
    },ensure_ascii=False,indent=2) + '\n')
    print(f'PPTX：{output}\n预览总图：{assets / "contact-sheet.jpg"}',flush=True)
    return output


def main():
    parser = argparse.ArgumentParser(description='从 Prim 在线展示截图导出 PPTX')
    parser.add_argument('--url', help='已有展示服务；省略则临时启动只读服务')
    parser.add_argument('--output', type=Path, default=ROOT / 'build/presentation/prim-class-presentation.pptx')
    parser.add_argument('--outline', type=Path, default=ROOT / 'docs/class-presentation.md')
    parser.add_argument('--main-only', action='store_true', help='仅导出主讲页，不包含附录')
    parser.add_argument('--presenter', default='')
    parser.add_argument('--course', default='')
    parser.add_argument('--date', default='')
    args = parser.parse_args()
    if args.url:
        capture(args.url, args.output.resolve(), args.main_only)
    else:
        with preview_server(args.outline.resolve(), args.presenter, args.course, args.date, default_runtime()) as (url, _):
            capture(url, args.output.resolve(), args.main_only)


if __name__ == '__main__':
    main()
